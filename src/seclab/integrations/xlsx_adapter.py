"""Adapter for the Vietnamese teacher Excel specification workbook."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from seclab.specification import LoadedSpecification, SourceDocument


class SpreadsheetSpecificationError(ValueError):
    """An actionable error tied to a workbook sheet or cell."""


def load_teacher_xlsx(path: str | Path) -> tuple[LoadedSpecification, dict[str, Any], dict[str, Any]]:
    """Convert the teacher workbook to one in-memory canonical schema document."""
    try:
        import openpyxl
    except ImportError as exc:
        raise SpreadsheetSpecificationError(
            "Thiếu thư viện đọc XLSX. Hãy cài ứng dụng bằng `pip install -e \".[gui,xlsx]\"`."
        ) from exc

    source = Path(path).expanduser().resolve(strict=True)
    raw_bytes = source.read_bytes()
    try:
        book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise SpreadsheetSpecificationError(f"Không mở được tệp XLSX: {exc}") from exc
    try:
        names = set(book.sheetnames)
        required = {"ThongTin", "YeuCauChucNang", "YeuCauPhiChucNang", "RangBuoc",
                    "MoHinhTrangThai", "ThuocTinhAnToan", "LoaiTruLanNhau", "HamSinhMa",
                    "KiemTraSoHoc", "TestCase"}
        optional_sheets = {"HuongDanChay", "SoTheoDoiBangChung", "MaTranTruyVet", "KiemToan"}
        sheet_labels = {
            "ThongTin": "Thông tin", "YeuCauChucNang": "Yêu cầu chức năng",
            "YeuCauPhiChucNang": "Yêu cầu phi chức năng", "RangBuoc": "Ràng buộc",
            "MoHinhTrangThai": "Mô hình trạng thái", "ThuocTinhAnToan": "Thuộc tính an toàn",
            "LoaiTruLanNhau": "Loại trừ lẫn nhau", "HamSinhMa": "Hàm sinh mã",
            "KiemTraSoHoc": "Kiểm tra số học", "TestCase": "Ca kiểm thử",
            "HuongDanChay": "Hướng dẫn chạy", "SoTheoDoiBangChung": "Theo dõi bằng chứng",
            "MaTranTruyVet": "Ma trận truy vết", "KiemToan": "Kiểm toán",
        }
        column_labels = {
            "ID": "Mã", "Mã kiểm thử": "Mã kiểm thử", "Mã yêu cầu": "Mã yêu cầu",
            "MoTa": "Mô tả", "DanhMuc": "Danh mục", "MucUuTien": "Mức ưu tiên",
            "TrangThai": "Trạng thái", "KhoiTao": "Khởi tạo", "Nhan": "Nhãn",
            "ChuyenDenCacTrangThai": "Chuyển tới trạng thái", "CongThucCTL": "Công thức CTL",
            "MenhDe1": "Mệnh đề 1", "MenhDe2": "Mệnh đề 2", "Ten": "Tên",
            "NgonNgu": "Ngôn ngữ", "Loai": "Loại", "ThamSo": "Tham số",
            "KichThuocBuffer": "Kích thước bộ đệm", "Loi": "Lỗi mô phỏng",
            "GiaTriBiMat1": "Giá trị thử nghiệm 1", "GiaTriBiMat2": "Giá trị thử nghiệm 2",
            "BienSoNguyen": "Biến số nguyên", "DieuKien": "Điều kiện",
            "BieuThucTran": "Biểu thức tràn", "YeuCauLienQuan": "Yêu cầu liên quan",
            "TieuDe": "Tiêu đề", "DieuKienTienQuyet": "Điều kiện tiên quyết",
            "CacBuoc": "Các bước", "KetQuaMongDoi": "Kết quả mong đợi",
            "Truong": "Trường", "GiaTri": "Giá trị",
            "Nhóm kiểm thử": "Nhóm kiểm thử", "Mã yêu cầu": "Mã yêu cầu",
            "Thuộc tính/điều kiện": "Thuộc tính/điều kiện", "Mục tiêu kiểm thử": "Mục tiêu kiểm thử",
            "Tiền điều kiện / dữ liệu đầu vào": "Tiền điều kiện / dữ liệu đầu vào",
            "Các bước thực hiện": "Các bước thực hiện",
            "Kết quả mong đợi / tiêu chí đạt": "Kết quả mong đợi / tiêu chí đạt",
            "Trạng thái thực thi": "Trạng thái thực thi", "Loại hiện vật kiểm thử": "Loại hiện vật kiểm thử",
            "Đường dẫn bằng chứng": "Đường dẫn bằng chứng", "Kết quả quan sát / ghi chú": "Kết quả quan sát / ghi chú",
        }
        # CTL proposition identifiers are ASCII tokens. Translating these
        # model labels breaks parsing and can make formulas refer to undeclared atoms.
        canonical_label_aliases = {
            "đã kết nối": "connected", "đã xác thực": "authenticated",
            "đang sạc": "charging", "safety_điều kiện cho phép_valid": "safety_gate_valid",
            "not_đã xác thực": "not_authenticated",
        }
        display_values = {
            "safety": "an toàn", "liveness": "tính sống", "functional": "chức năng",
            "security": "bảo mật", "performance": "hiệu năng", "other": "khác",
            "LOW": "THẤP", "MEDIUM": "TRUNG BÌNH", "HIGH": "CAO", "CRITICAL": "NGHIÊM TRỌNG",
            "C": "C", "python": "Python", "buffer_copy": "chép bộ đệm",
            "alloc_release": "cấp phát và giải phóng", "divide": "phép chia",
            "magic_value": "giá trị bí mật", "buffer-overflow": "tràn bộ đệm",
            "memory-leak": "rò rỉ bộ nhớ", "double-free": "giải phóng hai lần",
            "use-after-free": "truy cập sau khi giải phóng",
        }
        missing = sorted(required - names)
        if missing:
            raise SpreadsheetSpecificationError("Thiếu trang tính bắt buộc: " +
                ", ".join(sheet_labels.get(item, item) for item in missing))
        expected_headers = {
            "ThongTin": {"Truong", "GiaTri"},
            "YeuCauChucNang": {"ID", "MoTa", "DanhMuc", "MucUuTien"},
            "YeuCauPhiChucNang": {"ID", "MoTa", "MucUuTien"},
            "RangBuoc": {"ID", "MoTa", "MucUuTien"},
            "MoHinhTrangThai": {"TrangThai", "KhoiTao", "Nhan", "ChuyenDenCacTrangThai"},
            "ThuocTinhAnToan": {"CongThucCTL"},
            "LoaiTruLanNhau": {"MenhDe1", "MenhDe2"},
            "HamSinhMa": {"Ten", "NgonNgu", "Loai", "ThamSo", "KichThuocBuffer", "Loi",
                          "GiaTriBiMat1", "GiaTriBiMat2"},
            "KiemTraSoHoc": {"Ten", "BienSoNguyen", "DieuKien", "BieuThucTran"},
            "TestCase": {"Mã kiểm thử", "Nhóm kiểm thử", "Mã yêu cầu", "Thuộc tính/điều kiện",
                         "Mục tiêu kiểm thử", "Tiền điều kiện / dữ liệu đầu vào", "Các bước thực hiện",
                         "Kết quả mong đợi / tiêu chí đạt", "Trạng thái thực thi", "Loại hiện vật kiểm thử",
                         "Đường dẫn bằng chứng", "Kết quả quan sát / ghi chú"},
        }
        optional_headers = {"RangBuoc": {"YeuCauLienQuan", "Mã yêu cầu"},
            "KiemTraSoHoc": {"KetQuaSMTMongDoi", "BienThu/Boundary", "NegativeControl",
                            "RequirementIDs", "TrangThaiBangChung"}}

        def rows(sheet_name: str) -> list[tuple[int, dict[str, Any]]]:
            sheet = book[sheet_name]
            iterator = sheet.iter_rows(values_only=False)
            try:
                header_cells = next(iterator)
            except StopIteration:
                raise SpreadsheetSpecificationError(f"Trang tính '{sheet_labels[sheet_name]}' đang trống.")
            headers: dict[int, str] = {}
            for col, cell in enumerate(header_cells, 1):
                if cell.value is not None:
                    headers[col] = str(cell.value).strip()
            allowed_headers = expected_headers[sheet_name] | optional_headers.get(sheet_name, set())
            unknown_headers = set(headers.values()) - allowed_headers
            missing_headers = expected_headers[sheet_name] - set(headers.values())
            if unknown_headers or missing_headers:
                parts = []
                if missing_headers:
                    parts.append("thiếu cột: " + ", ".join(column_labels.get(item, item)
                                                              for item in sorted(missing_headers)))
                if unknown_headers:
                    parts.append("cột không thuộc biểu mẫu: " + ", ".join(
                        column_labels.get(item, item) for item in sorted(unknown_headers)))
                raise SpreadsheetSpecificationError(
                    f"Trang tính '{sheet_labels[sheet_name]}', dòng 1: " + "; ".join(parts) + ".")
            if len(set(headers.values())) != len(headers):
                raise SpreadsheetSpecificationError(
                    f"Trang tính '{sheet_labels[sheet_name]}', dòng 1: tên cột bị lặp.")
            result = []
            for row_num, cells in enumerate(iterator, 2):
                values: dict[str, Any] = {}
                for col, cell in enumerate(cells, 1):
                    value = cell.value
                    header = headers.get(col)
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        raise SpreadsheetSpecificationError(
                            f"Sheet '{sheet_name}', ô {cell.coordinate}: không hỗ trợ công thức Excel; "
                            "hãy thay bằng giá trị tĩnh."
                        )
                    if header and value is not None and str(value).strip() != "":
                        values[header] = value
                if values:
                    result.append((row_num, values))
            return result

        def display_value(value: Any) -> Any:
            if isinstance(value, str):
                return display_values.get(value, sheet_labels.get(value, value))
            return value

        def translated_source_rows(sheet_name: str) -> list[dict[str, Any]]:
            # Only the preview copy is translated; source bytes and canonical IDs
            # remain untouched for stable validation and integrity checking.
            translated_headers = {key: value for key, value in column_labels.items()}
            worksheet = book[sheet_name]
            values = worksheet.iter_rows(values_only=True)
            header = list(next(values, ()))
            result = []
            for row in values:
                if not any(item is not None for item in row):
                    continue
                result.append({translated_headers.get(str(key), str(key)): display_value(item)
                               for key, item in zip(header, row) if key is not None and item is not None})
            return result

        def translated_info() -> dict[str, Any]:
            translated = {}
            for _, row in rows("ThongTin"):
                key = str(row.get("Truong") or "").strip()
                value = row.get("GiaTri")
                if key:
                    translated[{"name": "Tên", "description": "Mô tả"}.get(key, key)] = value
            return translated

        def required_text(row: dict[str, Any], field: str, sheet: str, line: int) -> str:
            value = row.get(field)
            if value is None or not str(value).strip():
                raise SpreadsheetSpecificationError(f"Sheet '{sheet}', dòng {line}: thiếu '{field}'.")
            return str(value).strip()

        def split_list(value: Any) -> list[str]:
            if value is None:
                return []
            return [part.strip() for part in re.split(r"[,;\n]+", str(value)) if part.strip()]

        id_pattern = re.compile(r"^[A-Za-z][A-Za-z0-9_.:-]*$")
        invalid_ids: list[str] = []

        def check_id(sheet: str, line: int, field: str, value: str) -> None:
            if not id_pattern.fullmatch(value):
                invalid_ids.append(
                    f"{sheet_labels.get(sheet, sheet)}!{field}{line} = {value!r}")

        info_rows = rows("ThongTin")
        info = dict((str(k).strip(), v) for _, row in info_rows
                    if (k := row.get("Truong")) is not None for v in [row.get("GiaTri")])
        name = str(info.get("name") or source.stem).strip()
        spec_id = re.sub(r"[^A-Za-z0-9_.:-]+", "-", name).strip("-.")
        if not spec_id or not spec_id[0].isascii() or not spec_id[0].isalpha():
            spec_id = "XLSX-" + (spec_id or "specification")

        requirements: list[dict[str, Any]] = []
        requirement_ids_seen: set[str] = set()
        for sheet_name, category in (("YeuCauChucNang", "functional"),
                                     ("YeuCauPhiChucNang", "other")):
            for line, row in rows(sheet_name):
                requirement_id = required_text(row, "ID", sheet_name, line)
                check_id(sheet_name, line, "ID", requirement_id)
                if requirement_id in requirement_ids_seen:
                    raise SpreadsheetSpecificationError(
                        f"Sheet '{sheet_name}', dòng {line}: ID yêu cầu '{requirement_id}' bị trùng.")
                requirement_ids_seen.add(requirement_id)
                category_value = str(row.get("DanhMuc") or category).strip().casefold()
                item = {"id": requirement_id,
                        "description": required_text(row, "MoTa", sheet_name, line),
                        "type": category_value if category_value in
                            {"safety", "liveness", "functional", "security", "performance", "other"}
                            else category,
                        "metadata": {"sheet_nguon": sheet_name, "dong_nguon": line}}
                priority = row.get("MucUuTien")
                if priority:
                    priority_value = str(priority).strip().upper()
                    if priority_value not in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}:
                        raise SpreadsheetSpecificationError(
                            f"Sheet '{sheet_name}', dòng {line}: mức ưu tiên không hợp lệ '{priority_value}'.")
                    item["priority"] = priority_value
                requirements.append(item)
        req_ids = {item["id"] for item in requirements}

        constraint_ids_by_requirement: dict[str, list[str]] = {}
        constraints = []
        for line, row in rows("RangBuoc"):
            constraint_id = required_text(row, "ID", "RangBuoc", line)
            check_id("RangBuoc", line, "A", constraint_id)
            explicit_links = split_list(row.get("YeuCauLienQuan"))
            unknown_links = sorted(set(explicit_links) - req_ids)
            if unknown_links:
                raise SpreadsheetSpecificationError(
                    f"Sheet 'RangBuoc', dòng {line}: mã yêu cầu không tồn tại: {', '.join(unknown_links)}.")
            for requirement_id in explicit_links:
                constraint_ids_by_requirement.setdefault(requirement_id, []).append(constraint_id)
            if any(item["id"] == constraint_id for item in constraints):
                raise SpreadsheetSpecificationError(f"Sheet 'RangBuoc', dòng {line}: ID '{constraint_id}' bị trùng.")
            constraints.append({"id": constraint_id,
                "expression": required_text(row, "MoTa", "RangBuoc", line),
                "kind": "other", "description": required_text(row, "MoTa", "RangBuoc", line)})

        state_rows = rows("MoHinhTrangThai")
        states: list[dict[str, Any]] = []
        transitions: list[dict[str, Any]] = []
        initial_state = None
        for line, row in state_rows:
            state_id = required_text(row, "TrangThai", "MoHinhTrangThai", line)
            check_id("MoHinhTrangThai", line, "A", state_id)
            if any(item["id"] == state_id for item in states):
                raise SpreadsheetSpecificationError(
                    f"Sheet 'MoHinhTrangThai', dòng {line}: trạng thái '{state_id}' bị khai báo trùng.")
            initial = str(row.get("KhoiTao") or "").strip().casefold() in {"x", "yes", "true", "1", "có"}
            if initial:
                if initial_state is not None:
                    raise SpreadsheetSpecificationError(
                        f"Sheet 'MoHinhTrangThai': có nhiều trạng thái khởi tạo ('{initial_state}', '{state_id}').")
                initial_state = state_id
            states.append({"id": state_id, "initial": initial,
                           "labels": [canonical_label_aliases.get(label, label)
                                      for label in split_list(row.get("Nhan"))],
                           "description": "Trạng thái được nhập từ sheet MoHinhTrangThai."})
            for target in split_list(row.get("ChuyenDenCacTrangThai")):
                transitions.append({"id": f"TR-{len(transitions)+1:03d}", "source": state_id,
                    "target": target, "label": target,
                    })
        known_states = {state["id"] for state in states}
        for transition in transitions:
            if transition["target"] not in known_states:
                raise SpreadsheetSpecificationError(
                    f"Sheet 'MoHinhTrangThai': trạng thái '{transition['source']}' chuyển tới "
                    f"'{transition['target']}' nhưng trạng thái đích không được khai báo.")
        if not initial_state:
            raise SpreadsheetSpecificationError("Sheet 'MoHinhTrangThai': cần đánh dấu một trạng thái khởi tạo bằng 'x'.")

        properties = []
        ctl_atoms = {label for state in states for label in state["labels"]}
        state_ids = {state["id"] for state in states}
        for line, row in rows("ThuocTinhAnToan"):
            formula = required_text(row, "CongThucCTL", "ThuocTinhAnToan", line).strip()
            referenced = re.findall(r"req_([a-z0-9_]+)_violation", formula, re.I)
            linked = ["REQ-" + token.replace("_", "-").upper() for token in referenced]
            linked = [item for item in linked if item in req_ids]
            # The form mixes Vietnamese labels into CTL. Normalize known labels
            # back to the stable ASCII proposition vocabulary for the engine.
            for display_label, canonical_label in canonical_label_aliases.items():
                formula = re.sub(rf"(?<![A-Za-z0-9_.:-]){re.escape(display_label)}(?![A-Za-z0-9_.:-])",
                                 canonical_label, formula, flags=re.IGNORECASE)
            try:
                from seclab.formal.ctl.ast import atomic_propositions
                from seclab.formal.ctl.parser import parse_ctl
                atoms = atomic_propositions(parse_ctl(formula))
            except Exception as exc:
                raise SpreadsheetSpecificationError(
                    f"Sheet 'ThuocTinhAnToan', ô A{line}: công thức CTL không hợp lệ: {exc}") from exc
            unknown_atoms = sorted(atom for atom in atoms
                                   if atom not in ctl_atoms
                                   and atom not in state_ids
                                   and not re.fullmatch(r"req_[a-z0-9_]+_violation", atom, re.I))
            if unknown_atoms:
                raise SpreadsheetSpecificationError(
                    f"Sheet 'ThuocTinhAnToan', ô A{line}: các nhãn chưa khai báo trong cột 'Nhan' "
                    f"hoặc chưa có trạng thái tương ứng: {', '.join(unknown_atoms)}.")
            properties.append({"id": f"CTL-P-{len(properties)+1:03d}", "logic": "CTL",
                "formula": formula, "requirement_ids": linked,
                "description": "Công thức CTL nhập từ sheet ThuocTinhAnToan.",
                })

        functions = []
        function_ids_seen: set[str] = set()
        for line, row in rows("HamSinhMa"):
            fname = required_text(row, "Ten", "HamSinhMa", line)
            fragment = re.sub(r"[^A-Za-z0-9_.:-]+", "-", fname).strip("-.")
            if not fragment or not fragment[0].isalpha():
                fragment = "FN-" + (fragment or "function")
            function_id = "FN-" + fragment
            if function_id in function_ids_seen:
                raise SpreadsheetSpecificationError(
                    f"Sheet 'HamSinhMa', dòng {line}: tên hàm '{fname}' tạo ID bị trùng '{function_id}'.")
            function_ids_seen.add(function_id)
            functions.append({"id": function_id, "name": fname,
                "language": str(row.get("NgonNgu") or "unknown").strip().casefold(),
                "description": f"Khai báo trong bảng tính (loại={row.get('Loai')}, tham số={row.get('ThamSo')}, kích thước={row.get('KichThuocBuffer')}, lỗi mô phỏng={row.get('Loi')}). Không suy diễn thành mã nguồn thực thi.",
                })

        tests = []
        seen_test_ids: set[str] = set()
        for line, row in rows("TestCase"):
            test_id = required_text(row, "Mã kiểm thử", "TestCase", line)
            check_id("TestCase", line, "A", test_id)
            if test_id in seen_test_ids:
                raise SpreadsheetSpecificationError(f"Sheet 'TestCase', dòng {line}: ID '{test_id}' bị trùng.")
            seen_test_ids.add(test_id)
            declared_refs = split_list(row.get("Mã yêu cầu"))
            links = [item for item in declared_refs if item in req_ids]
            # The teacher form also stores evidence-scope tags in this column,
            # such as PROBE_ONLY and TOOL_CAPABILITY_EVIDENCE. Those are not
            # requirement IDs and must not be passed to the canonical linker.
            unknown_links = sorted({item for item in declared_refs
                                    if item not in req_ids and not re.fullmatch(
                                        r"[A-Z][A-Z0-9_]*", item)})
            if unknown_links:
                raise SpreadsheetSpecificationError(
                    f"Trang tính 'Ca kiểm thử', dòng {line}: mã yêu cầu không tồn tại: {', '.join(unknown_links)}.")
            test = {"id": test_id, "kind": "system", "requirement_ids": links,
                "description": "\n".join(f"{key}: {row[key]}" for key in
                    ("Nhóm kiểm thử", "Thuộc tính/điều kiện", "Mục tiêu kiểm thử",
                     "Tiền điều kiện / dữ liệu đầu vào", "Các bước thực hiện",
                     "Kết quả mong đợi / tiêu chí đạt", "Trạng thái thực thi",
                     "Loại hiện vật kiểm thử", "Đường dẫn bằng chứng", "Kết quả quan sát / ghi chú")
                    if row.get(key))}
            tests.append(test)

        if invalid_ids:
            raise SpreadsheetSpecificationError(
                "Các mã định danh phải bắt đầu bằng chữ cái ASCII và chỉ chứa chữ cái, chữ số, '_', '.', ':', '-'. "
                "Hãy sửa các ô: " + "; ".join(invalid_ids))

        requirement_properties: dict[str, list[str]] = {}
        for prop in properties:
            for requirement_id in prop["requirement_ids"]:
                requirement_properties.setdefault(requirement_id, []).append(prop["id"])
        requirement_tests: dict[str, list[str]] = {}
        for test in tests:
            for requirement_id in test["requirement_ids"]:
                requirement_tests.setdefault(requirement_id, []).append(test["id"])
        for requirement in requirements:
            requirement["constraint_ids"] = constraint_ids_by_requirement.get(requirement["id"], [])
            requirement["property_ids"] = requirement_properties.get(requirement["id"], [])
            requirement["test_ids"] = requirement_tests.get(requirement["id"], [])

        canonical = {"schema_version": "2.0", "id": spec_id, "name": name,
            "description": str(info.get("description") or "Đặc tả nhập từ workbook XLSX theo form giáo viên."),
            "initial_state": initial_state, "requirements": requirements, "variables": [],
            "constraints": constraints, "states": states, "transitions": transitions,
            "properties": properties, "functions": functions, "tests": tests,
            "metadata": {"dinh_dang_nguon": "bieu_mau_xlsx_giao_vien",
                "cac_sheet_nguon": sorted(required),
                "cac_cap_loai_tru": [{"menh_de_1": row.get("MenhDe1"), "menh_de_2": row.get("MenhDe2")}
                                       for _, row in rows("LoaiTruLanNhau")],
                "cac_kiem_tra_so_hoc": [{"ten": row.get("Ten"), "bien_so_nguyen": row.get("BienSoNguyen"),
                    "dieu_kien": row.get("DieuKien"), "bieu_thuc_tran": row.get("BieuThucTran")}
                    for _, row in rows("KiemTraSoHoc")],
                "ghi_chu_nhap_xlsx": ["Biểu thức số học và ca kiểm thử dạng mô tả được lưu làm dữ liệu nguồn, không tự chạy.",
                    "Sheet HamSinhMa là khai báo/probe; không tạo mã triển khai hoặc xác nhận lỗi trong mã nguồn sản phẩm."]}}
        digest = hashlib.sha256(raw_bytes).hexdigest()
        canonical["metadata"]["nguon_goc"] = {"tep_nguon": source.name, "ma_bam_sha256": digest}
        loaded = LoadedSpecification((SourceDocument(source, canonical, digest),))
        preview_sheets = required | (optional_sheets & names)
        raw = {"sheets": {sheet_labels[sheet]: translated_source_rows(sheet)
                          for sheet in sorted(preview_sheets)}}
        raw["sheets"][sheet_labels["ThongTin"]] = translated_info()
        provenance = {"format": "Đặc tả giáo viên XLSX", "adapter": "Bộ chuyển đổi XLSX",
                      "original_sha256": digest, "source_format": "XLSX",
                      "normalization_status": "IN_MEMORY", "sheet_count": len(preview_sheets)}
        return loaded, provenance, raw
    finally:
        book.close()
